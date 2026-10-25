import os,time,secrets,json,base64,asyncio,io,zipfile
from contextlib import asynccontextmanager
from typing import Literal
from fastapi import FastAPI,HTTPException,Header,UploadFile,File,Depends,Response
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel,Field
from . import core,store,providers,media,evidence

class Config(BaseModel):
 candidate_name:str=Field(min_length=1,max_length=120)
 job_title:str=Field(min_length=1,max_length=160)
 job_description:str=Field(min_length=10,max_length=20000)
 resume_text:str=Field(default='',max_length=30000)
 skills:list[str]=Field(min_length=1,max_length=30)
 experience_level:str='mid-level'
 difficulty:Literal['beginner','intermediate','advanced']='intermediate'
 question_count:int|None=Field(default=10,ge=1,le=30)
 duration_minutes:int=Field(default=25,ge=2,le=120)
 answer_seconds:int=Field(default=120,ge=15,le=600)
 custom_questions:list[str]=Field(default_factory=list,max_length=30)
 categories:list[Literal['technical','scenario','experience','behavioral']]=Field(default=['technical','scenario','experience','behavioral'],min_length=1)
 followups:bool=True
 candidate_profile:dict|None=None
 skill_alignment:dict|None=None
class AnalyzeProfileRequest(BaseModel):
 resume_text:str=Field(default='',max_length=35000)
 job_description:str=Field(default='',max_length=20000)
 skills:list[str]=Field(default_factory=list)
class QuestionUpdate(BaseModel):
 text:str=Field(min_length=5,max_length=2000)
 category:Literal['technical','scenario','experience','behavioral','recruiter']|None=None
 skill:str|None=None
class Consent(BaseModel):
 accepted:bool
 policy_version:Literal['2026-10-25']='2026-10-25'
class Action(BaseModel):
 action:Literal['answer','partial','repeat','request_skip','cancel_skip','skip','finish']
 text:str=Field(default='',max_length=20000)
 slot_id:str
 request_id:str=Field(min_length=8,max_length=100)
class Override(BaseModel):
 target_id:str
 reason:str=Field(min_length=8,max_length=2000)
 score:int|None=Field(default=None,ge=1,le=5)
 disposition:Literal['confirmed','dismissed','needs_review']|None=None
class Pair(BaseModel):
 nonce:str=Field(min_length=32,max_length=100)
 accepted:bool
class Event(BaseModel):
 role:Literal['primary','secondary']
 type:str=Field(max_length=100)
 severity:Literal['OBSERVED_EVENT','NEEDS_REVIEW','NEEDS_REVIEW_HIGH','UNVERIFIED']
 jpeg_b64:str=Field(max_length=1500000)
 algorithm:str=Field(max_length=100)
 frame_count:int=Field(ge=1)
 start:float
 end:float

@asynccontextmanager
async def lifespan(app):
 for k in ('RECRUITER_TOKEN','WORKER_TOKEN','SIGNING_KEY','EVIDENCE_KEY'):
  if len(os.getenv(k,''))<24:raise RuntimeError(k+' missing/too short. Run scripts/setup.py')
 store.init();yield
app=FastAPI(title='HireNeuron Verity',version='0.1.0',lifespan=lifespan)
app.add_middleware(CORSMiddleware,allow_origins=os.getenv('WEB_ORIGINS','http://localhost:3000').split(','),allow_methods=['GET','POST','DELETE'],allow_headers=['Authorization','Content-Type'])

def auth(authorization:str=Header(default='')):
 if not authorization.startswith('Bearer '):raise HTTPException(401,'Bearer token required')
 t=authorization[7:]
 if secrets.compare_digest(t,os.getenv('RECRUITER_TOKEN','')):return {'role':'recruiter','sid':None}
 if secrets.compare_digest(t,os.getenv('WORKER_TOKEN','')):return {'role':'worker','sid':None}
 with store.connect() as c:
  row=c.execute('SELECT * FROM credentials WHERE hash=? AND expires>?',(core.digest(t),time.time())).fetchone()
 if not row:raise HTTPException(401,'Invalid or expired credential')
 return dict(row)
def allow(a,sid,roles):
 if a['role'] not in roles or (a['sid'] and a['sid']!=sid):raise HTTPException(403,'Access denied')
def mint(c,sid,role,ttl=7200):
 t=secrets.token_urlsafe(32);c.execute('INSERT INTO credentials VALUES (?,?,?,?)',(core.digest(t),sid,role,time.time()+ttl));return t

def public(s,a):
 if a['role'] in ('recruiter','worker'):return s
 q=core.active(s)
 return {k:s[k] for k in ('id','state','current','revision','started_at','question_started_at','phase','pending_skip')}|{
  'candidate_name':s['config']['candidate_name'],'job_title':s['config']['job_title'],
  'question_count':len(s['questions']),'duration_minutes':s['config']['duration_minutes'],
  'answer_seconds':s['config']['answer_seconds'],
  'question':{k:q[k] for k in ('id','slot','text','followup','status','spans')} if q else None}
@app.exception_handler(KeyError)
async def missing(request,exc):
 from fastapi.responses import JSONResponse
 return JSONResponse(status_code=404,content={'detail':'Not found'})
@app.exception_handler(ValueError)
async def invalid(request,exc):
 from fastapi.responses import JSONResponse
 return JSONResponse(status_code=409,content={'detail':str(exc)})
@app.get('/health')
def health():return {'status':'ok','mode':'prototype','model_configured':bool(os.getenv('GEMINI_MODEL')),'media_configured':bool(os.getenv('LIVEKIT_URL'))}
@app.post('/sessions')
def create(cfg:Config,a=Depends(auth)):
 allow(a,None,['recruiter']);s=core.new_session(cfg.model_dump())
 with store.tx() as c:
  store.put(c,s);t=mint(c,s['id'],'candidate',86400)
 return {'session':s,'candidate_token':t}
@app.get('/sessions')
def listing(a=Depends(auth)):
 allow(a,None,['recruiter','worker'])
 with store.connect() as c:rows=c.execute('SELECT data FROM sessions').fetchall()
 return [dict(id=s['id'],candidate=s['config']['candidate_name'],role=s['config']['job_title'],state=s['state'],created_at=s['created_at']) for s in [json.loads(r[0]) for r in rows]]
@app.get('/sessions/{sid}')
def get(sid:str,a=Depends(auth)):
 allow(a,sid,['recruiter','candidate','worker'])
 with store.tx() as c:
  s=store.get(c,sid);before=core.canonical(s);core.tick(s,time.time())
  if before!=core.canonical(s):store.put(c,s)
 return public(s,a)
@app.post('/sessions/{sid}/questions')
async def generate_questions(sid:str,a=Depends(auth)):
 allow(a,sid,['recruiter'])
 with store.connect() as c:s=store.get(c,sid)
 if s['state']!='BLUEPRINT_FROZEN':raise HTTPException(409,'Cannot rewrite a started interview')
 try:s=await providers.questions(s)
 except Exception as e:raise HTTPException(503,'Question provider unavailable/invalid response; frozen templates retained') from e
 with store.tx() as c:
  latest=store.get(c,sid)
  if latest['state']!='BLUEPRINT_FROZEN':raise HTTPException(409,'State changed; retry')
  s['coverage']=latest.get('coverage',{})
  store.put(c,s)
 return s
@app.patch('/sessions/{sid}/questions/{qid}')
def update_question(sid:str,qid:str,body:QuestionUpdate,a=Depends(auth)):
 allow(a,sid,['recruiter'])
 with store.tx() as c:
  s=store.get(c,sid)
  if s['state']!='BLUEPRINT_FROZEN':raise HTTPException(409,'Cannot modify blueprint after interview started')
  q=next((x for x in s['questions'] if x['id']==qid),None)
  if not q:raise HTTPException(404,'Question not found')
  q['text']=body.text.strip();q['locked']=True
  if body.category:q['category']=body.category
  if body.skill:q['skill']=body.skill.strip()
  store.put(c,s)
 return {'question':q}
@app.post('/documents/extract')
async def extract(file:UploadFile=File(...),a=Depends(auth)):
 allow(a,None,['recruiter']);data=await file.read(5*1024*1024+1)
 if len(data)>5*1024*1024:raise HTTPException(413,'Maximum 5 MB')
 name=(file.filename or '').lower()
 try:
  if name.endswith('.pdf') and data.startswith(b'%PDF'):
   from pypdf import PdfReader
   reader=PdfReader(io.BytesIO(data))
   if len(reader.pages)>30:raise ValueError('Maximum 30 pages')
   text='\n'.join(p.extract_text() or '' for p in reader.pages)
  elif name.endswith('.docx') and data.startswith(b'PK'):
   with zipfile.ZipFile(io.BytesIO(data)) as z:
    if sum(f.file_size for f in z.infolist())>25*1024*1024:raise ValueError('Expanded document too large')
   from docx import Document
   doc=Document(io.BytesIO(data));text='\n'.join(p.text for p in doc.paragraphs)
   text+='\n'+'\n'.join(' | '.join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)
  else:raise ValueError('Use PDF or DOCX')
 except Exception as e:raise HTTPException(422,'Unreadable file, encrypted document, or unsupported content') from e
 return {'text':text[:30000],'needs_ocr':len(text.strip())<40,'note':'Text extraction only. Review and correct before generating questions.'}
@app.post('/documents/analyze-profile')
async def analyze_profile_endpoint(b:AnalyzeProfileRequest,a=Depends(auth)):
 allow(a,None,['recruiter'])
 return await providers.analyze_profile(b.resume_text,b.job_description,b.skills)
@app.post('/sessions/{sid}/consent')
def consent(sid:str,body:Consent,a=Depends(auth)):
 allow(a,sid,['candidate'])
 if not body.accepted:raise HTTPException(409,'Consent declined; do not start capture')
 with store.tx() as c:
  s=store.get(c,sid)
  if not any(x['role']=='candidate' for x in s['consents']):s['consents'].append({'role':'candidate','policy':body.policy_version,'at':time.time()})
  store.put(c,s)
 return {'accepted':True}
@app.post('/sessions/{sid}/start')
def start(sid:str,a=Depends(auth)):
 allow(a,sid,['candidate'])
 with store.tx() as c:s=store.get(c,sid);core.start(s,time.time());store.put(c,s)
 return public(s,a)
@app.post('/sessions/{sid}/action')
def action(sid:str,body:Action,a=Depends(auth)):
 allow(a,sid,['candidate','worker'])
 with store.tx() as c:
  s=store.get(c,sid)
  if body.request_id in s['processed']:
   saved=s['processed'][body.request_id]
   if isinstance(saved,dict) and saved['payload']!=core.digest(body.model_dump_json()):raise HTTPException(409,'Request ID reused with a different payload')
   return {'state':public(s,a),'message':'Already applied'}
  message=core.command(s,body.action,body.text,time.time(),body.slot_id)
  if a['role']=='worker' and body.action in ('partial','answer') and body.text.strip():
   for q in s['questions']:
    if q['id']==body.slot_id and q['spans']:q['spans'][-1]['source']='server_stt'
  s['processed'][body.request_id]={'at':time.time(),'payload':core.digest(body.model_dump_json())};store.put(c,s)
 return {'state':public(s,a),'message':message}
@app.post('/sessions/{sid}/followup')
async def followup(sid:str,body:Action,a=Depends(auth)):
 allow(a,sid,['candidate','worker'])
 with store.connect() as c:s=store.get(c,sid)
 core.tick(s,time.time());q=core.active(s)
 if s['state']!='LISTENING' or not q or q['id']!=body.slot_id or q['followup'] or not s['config']['followups']:return {'needed':False}
 remaining=s['config']['duration_minutes']*60-(time.time()-s['started_at'])
 if remaining<60 or time.time()-s['question_started_at']>s['config']['answer_seconds']-30:return {'needed':False}
 try:r=await providers.generate('Ask at most one short contextual follow-up only when a specific answer claim needs probing. Do not answer for candidate.',{'question':q['text'],'answer':body.text},providers.Followup)
 except Exception:return {'needed':False}
 if r['needed'] and r['text'].strip():
  with store.tx() as c:
   current=store.get(c,sid);core.tick(current,time.time());cq=core.active(current)
   if not cq or cq['id']!=q['id'] or cq['followup'] or cq['spans']!=q['spans']:return {'needed':False}
   if current['state']!='LISTENING':return {'needed':False}
   cq['spans'].append({'id':core.uid(),'text':body.text,'phase':'primary','offset':time.time()-current['started_at'],'source':'candidate_text'})
   cq['followup']=r['text'];current['phase']='followup';store.put(c,current)
 return r
@app.post('/sessions/{sid}/pair')
def pair(sid:str,a=Depends(auth)):
 allow(a,sid,['candidate'])
 with store.tx() as c:
  s=store.get(c,sid)
  if s['state'] not in ('BLUEPRINT_FROZEN','LISTENING'):raise HTTPException(409,'Consent/active session required')
  c.execute('DELETE FROM pairing WHERE sid=?',(sid,))
  nonce=secrets.token_urlsafe(32);c.execute('INSERT INTO pairing(hash,sid,expires) VALUES (?,?,?)',(core.digest(nonce),sid,time.time()+180))
 return {'nonce':nonce,'expires_in':180}
@app.post('/pair/redeem')
def redeem(body:Pair):
 if not body.accepted:raise HTTPException(409,'Separate phone consent required')
 with store.tx() as c:
  row=c.execute('SELECT * FROM pairing WHERE hash=? AND used=0 AND expires>?',(core.digest(body.nonce),time.time())).fetchone()
  if not row:raise HTTPException(410,'Pairing link used or expired')
  s=store.get(c,row['sid'])
  if s['state'] not in ('BLUEPRINT_FROZEN','LISTENING'):raise HTTPException(409,'Inactive interview')
  c.execute('UPDATE pairing SET used=1 WHERE hash=?',(row['hash'],));c.execute('DELETE FROM credentials WHERE sid=? AND role=?',(s['id'],'phone'))
  t=mint(c,s['id'],'phone');s['consents'].append({'role':'phone','policy':'2026-10-25','at':time.time()});store.put(c,s)
 return {'session_id':s['id'],'token':t}
@app.post('/sessions/{sid}/media')
def media_token(sid:str,a=Depends(auth)):
 allow(a,sid,['candidate','phone','recruiter'])
 with store.connect() as c:s=store.get(c,sid)
 if a['role'] not in ('recruiter','candidate') and not any(x['role']==a['role'] for x in s['consents']):raise HTTPException(403,'Consent missing')
 if s['state'] not in ('BLUEPRINT_FROZEN','LISTENING'):raise HTTPException(409,'Session inactive')
 if not os.getenv('LIVEKIT_URL'):raise HTTPException(503,'LiveKit not configured; text mode available')
 return {'url':os.environ['LIVEKIT_URL'],'token':media.token(sid,a['role'])}
@app.get('/sessions/{sid}/events')
def list_events(sid:str,a=Depends(auth)):
 allow(a,sid,['candidate','recruiter'])
 with store.connect() as c:s=store.get(c,sid)
 return s.get('events',[])

@app.post('/sessions/{sid}/events')
def event(sid:str,body:Event,a=Depends(auth)):
 allow(a,sid,['worker'])
 from PIL import Image
 try:
  data=base64.b64decode(body.jpeg_b64,validate=True)
  im=Image.open(io.BytesIO(data));im.verify()
 except Exception as e:raise HTTPException(422,'Invalid snapshot') from e
 with store.tx() as c:
  s=store.get(c,sid)
  if s['state'] not in ('BLUEPRINT_FROZEN','LISTENING'):raise HTTPException(409,'Capture session inactive')
  role='candidate' if body.role=='primary' else 'phone'
  if not any(x['role']==role for x in s['consents']):raise HTTPException(403,'Missing camera consent')
  meta=body.model_dump(exclude={'jpeg_b64'});meta['received_at']=time.time();meta['offset']=max(0,time.time()-(s['started_at'] or time.time()))
  record=evidence.persist(s,data,meta);store.put(c,s)
 return record
@app.get('/sessions/{sid}/evidence/{eid}')
def snapshot(sid:str,eid:str,a=Depends(auth)):
 allow(a,sid,['recruiter'])
 with store.connect() as c:s=store.get(c,sid)
 if not any(e['id']==eid for e in s['events']):raise HTTPException(404,'Not found')
 return Response(evidence.read(eid),media_type='image/jpeg',headers={'Cache-Control':'no-store'})
@app.post('/sessions/{sid}/evaluate')
def enqueue(sid:str,a=Depends(auth)):
 allow(a,sid,['recruiter'])
 with store.tx() as c:
  s=store.get(c,sid)
  if s['state'] not in ('COMPLETED','REVIEW_READY'):raise HTTPException(409,'Complete interview first')
  c.execute("INSERT INTO jobs VALUES (?,'queued',?) ON CONFLICT(sid) DO UPDATE SET state='queued',updated=excluded.updated WHERE jobs.state!='running'",(sid,time.time()))
 return {'queued':True}
@app.post('/sessions/{sid}/override')
def override(sid:str,b:Override,a=Depends(auth)):
 allow(a,sid,['recruiter'])
 with store.tx() as c:
  s=store.get(c,sid);target=next((q for q in s['questions'] if q['id']==b.target_id),None)
  ev=next((e for e in s['events'] if e['id']==b.target_id),None)
  if not target and not ev:raise HTTPException(404,'Target not found')
  if (target and b.score is None) or (ev and b.disposition is None):raise HTTPException(422,'Score or event disposition required')
  s['audit'].append(dict(id=core.uid(),reviewer='prototype-recruiter',at=time.time(),original=target['scores'] if target else ev['severity'],**b.model_dump()));store.put(c,s)
 return {'saved':True}
@app.delete('/sessions/{sid}')
def delete(sid:str,a=Depends(auth)):
 allow(a,sid,['recruiter'])
 from pathlib import Path
 with store.tx() as c:
  s=store.get(c,sid)
  if s['state']=='LISTENING':raise HTTPException(409,'Finish active session before deletion')
  for e in s['events']:(Path(os.getenv('EVIDENCE_PATH','data/evidence'))/(e['id']+'.enc')).unlink(missing_ok=True)
  rec_dir=Path(os.getenv('RECORDINGS_PATH','data/recordings'))
  for ext in ('.webm','.mp4'):(rec_dir/(sid+ext)).unlink(missing_ok=True)
  for table in ('credentials','pairing','jobs'):c.execute(f'DELETE FROM {table} WHERE sid=?',(sid,))
  c.execute('DELETE FROM sessions WHERE id=?',(sid,))
 return {'deleted':True}

@app.post('/sessions/{sid}/recording/upload')
async def upload_recording(sid:str,file:UploadFile=File(...),a=Depends(auth)):
 allow(a,sid,['candidate','recruiter','worker'])
 from pathlib import Path
 with store.tx() as c:
  s=store.get(c,sid)
  if not any(x['role'] in ('candidate_recording','candidate') for x in s['consents']):
   raise HTTPException(403,'Recording consent required')
  data=await file.read(150*1024*1024+1)
  if len(data)>150*1024*1024:raise HTTPException(413,'Maximum recording size 150MB')
  import hashlib
  sha=hashlib.sha256(data).hexdigest()
  rec_dir=Path(os.getenv('RECORDINGS_PATH','data/recordings'))
  rec_dir.mkdir(parents=True,exist_ok=True)
  ext='.webm' if (file.filename or '').endswith('.webm') else '.mp4'
  rec_path=rec_dir/(sid+ext)
  rec_path.write_bytes(data)
  dur=max(0,round(time.time()-(s['started_at'] or time.time()),1))
  s['recording']={'id':core.uid(),'filename':sid+ext,'sha256':sha,'size':len(data),'duration_seconds':dur,'created_at':time.time(),'status':'finalized','provenance':'Consented candidate laptop browser MediaRecorder capture (primary camera + microphone)'}
  store.put(c,s)
 return {'status':'uploaded','sha256':sha,'size':len(data)}

@app.get('/sessions/{sid}/recording')
def get_recording(sid:str,a=Depends(auth)):
 allow(a,sid,['recruiter'])
 from pathlib import Path
 with store.connect() as c:s=store.get(c,sid)
 rec=s.get('recording')
 if not rec:raise HTTPException(404,'No recording available')
 rec_dir=Path(os.getenv('RECORDINGS_PATH','data/recordings'))
 p=rec_dir/rec['filename']
 if not p.exists():raise HTTPException(404,'Recording file not found')
 media_type='video/webm' if p.suffix=='.webm' else 'video/mp4'
 from fastapi.responses import FileResponse
 return FileResponse(p,media_type=media_type)

@app.delete('/sessions/{sid}/recording')
def delete_recording(sid:str,a=Depends(auth)):
 allow(a,sid,['recruiter'])
 from pathlib import Path
 with store.tx() as c:
  s=store.get(c,sid)
  rec=s.get('recording')
  if rec:
   rec_dir=Path(os.getenv('RECORDINGS_PATH','data/recordings'))
   (rec_dir/rec['filename']).unlink(missing_ok=True)
   s['recording']=None
   store.put(c,s)
 return {'deleted':True}

class Coverage(BaseModel):
 worker:Literal['integrity','voice']
 frames:dict[str,float]=Field(default_factory=dict)
 algorithm:str=Field(max_length=100)
@app.post('/sessions/{sid}/coverage')
def coverage(sid:str,b:Coverage,a=Depends(auth)):
 allow(a,sid,['worker'])
 with store.tx() as c:
  s=store.get(c,sid);s['coverage'][b.worker]={'received_at':time.time(),**b.model_dump()};store.put(c,s)
 return {'ok':True}

@app.get('/sessions/{sid}/phone-status')
def phone_status(sid:str,a=Depends(auth)):
 allow(a,sid,['phone'])
 with store.connect() as c:s=store.get(c,sid)
 return {'active':s['state'] in ('BLUEPRINT_FROZEN','LISTENING')}
